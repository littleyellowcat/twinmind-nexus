from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "TwinMind_Nexus_孪生智枢_项目面试准备手册_第一版.docx"


BLUE = RGBColor(31, 78, 121)
DARK = RGBColor(33, 37, 41)
GRAY = RGBColor(88, 88, 88)
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F5F6F8"
LIGHT_YELLOW = "FFF6D8"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    set_run_font(run)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, east_asia: str = "Microsoft YaHei") -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.color.rgb = DARK


def add_para(doc: Document, text: str = "", style: str | None = None, bold_prefix: str | None = None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        set_run_font(r1)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.08
        run = p.add_run(item)
        set_run_font(run)
        run.font.size = Pt(10)


def add_numbers(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.08
        run = p.add_run(item)
        set_run_font(run)
        run.font.size = Pt(10)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(level=level)
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_run_font(run)
    run.font.color.rgb = BLUE if level <= 2 else RGBColor(67, 67, 67)
    run.font.bold = True
    run.font.size = Pt(16 if level == 1 else 13 if level == 2 else 11.5)


def add_callout(doc: Document, title: str, body: str, fill: str = LIGHT_YELLOW) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Cm(17.0)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_run_font(r)
    r.bold = True
    r.font.size = Pt(10)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_run_font(r2)
    r2.font.size = Pt(9.5)
    doc.add_paragraph()


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    if widths:
        for i, width in enumerate(widths):
            table.columns[i].width = Cm(width)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_text(cell, header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text)
            if i == 0:
                cells[i].paragraphs[0].runs[0].bold = True
    doc.add_paragraph()


def add_code_like(doc: Document, title: str, lines: list[str]) -> None:
    add_para(doc, title, bold_prefix=title.split("：")[0] + "：" if "：" in title else None)
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Cm(17.0)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for idx, line in enumerate(lines):
        if idx:
            p.add_run("\n")
        r = p.add_run(line)
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.size = Pt(8.8)
        r.font.color.rgb = RGBColor(30, 30, 30)
    doc.add_paragraph()


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for style_name in ["List Bullet", "List Number"]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10)
        style.paragraph_format.space_after = Pt(3)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("项目面试准备手册 · TwinMind Nexus（孪生智枢）")
    set_run_font(run)
    run.font.size = Pt(8.5)
    run.font.color.rgb = GRAY

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("A4 竖版打印稿 · 第一版")
    set_run_font(run)
    run.font.size = Pt(8.5)
    run.font.color.rgb = GRAY
    return doc


def build_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(36)
    r = p.add_run("项目面试准备手册")
    set_run_font(r)
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = BLUE

    p = doc.add_paragraph()
    r = p.add_run("TwinMind Nexus（孪生智枢）")
    set_run_font(r)
    r.font.size = Pt(15)
    r.font.bold = True

    add_para(
        doc,
        "定位：给小白准备的项目复盘稿，帮助你在面试中讲清楚 TwinMind Nexus 的业务背景、操作流程、技术原理、代码链路、Agent 设计、Graph RAG 设计、评测指标和常见追问。",
    )
    add_callout(
        doc,
        "阅读方式",
        "先背每个项目的 1 分钟介绍，再理解“用户一次操作会经过哪些函数和服务”，最后重点准备高优先级面试题。遇到英文方法名时，旁边都有中文解释，方便你把代码实现讲成人话。",
        LIGHT_BLUE,
    )
    add_table(
        doc,
        ["项目定位", "核心能力", "面试重点"],
        [
            [
                "多模态项目数字孪生 + 自主 Agent 图谱 RAG 平台",
                "上传完整项目 ZIP 后，自动生成实体、关系、证据卡、知识图谱、Hybrid RAG 索引和 Agent 报告。",
                "Plan-and-Execute、bounded ReAct、Agent 工具调用、知识图谱、证据链、AgentEval。",
            ],
        ],
        [4.8, 6.4, 5.8],
    )
    doc.add_page_break()


def build_project_one(doc: Document) -> None:
    add_heading(doc, "项目一：Modular RAG MCP Server", 1)
    add_para(
        doc,
        "这个项目可以理解为“企业私有知识库问答服务”。用户把 PDF、Markdown、技术文档放进系统，系统把文档拆成小块、向量化、建关键词索引；当 AI 助手提问时，通过 MCP 工具调用检索服务，返回带引用来源的答案。",
    )

    add_heading(doc, "1. 业务逻辑怎么讲", 2)
    add_table(
        doc,
        ["角色", "实际需求", "系统怎么解决"],
        [
            ["业务用户", "想问：某个配置怎么写、某个流程在哪里说明。", "输入自然语言问题，系统返回 answer（回答）和 citations（引用来源）。"],
            ["企业/团队", "文档多、分散、关键词搜不准。", "把 PDF/Markdown 统一摄取到 RAG 知识库，支持语义检索和关键词检索。"],
            ["AI 助手", "需要调用外部私有知识。", "通过 MCP 暴露 `query_knowledge_hub` 等工具，让 Copilot/Claude 类助手可以标准化调用。"],
        ],
        [3.1, 6.2, 7.7],
    )

    add_heading(doc, "2. 用户一次操作的完整流程", 2)
    add_numbers(
        doc,
        [
            "管理员上传知识文档，例如 PDF、Markdown、README、内部规范。",
            "Ingestion Pipeline（摄取管线）把文档解析为 Markdown 或纯文本。",
            "DocumentChunker（文档分块器）把长文档拆成 chunk（小文本块），每个 chunk 带 source_path、页码、标题等 metadata。",
            "Embedding 模型把 chunk 转成 dense vector（稠密向量），写入 ChromaDB。",
            "SparseEncoder 统计词频，BM25Indexer 建立 sparse index（关键词索引）。",
            "用户提问时，`query_knowledge_hub` 调用 HybridSearch，同时做向量检索和 BM25 检索。",
            "RRFFusion 把两路结果融合；如果启用 Cross-Encoder，再做 rerank（重排序）。",
            "ResponseBuilder 生成 answer（回答）和 citations（证据引用），返回给 MCP 客户端或前端。",
        ],
    )

    add_code_like(
        doc,
        "例子：用户问“Azure OpenAI 的 api_key 在哪里配置？”",
        [
            "输入 question（问题）: Azure OpenAI 的 api_key 在哪里配置？",
            "Dense Retrieval（向量检索）: 找到语义上接近“Azure / API Key / 配置”的 chunk。",
            "BM25（关键词检索）: 精确匹配 azure、api_key、settings.yaml 等词。",
            "RRF（融合排序）: 如果某个 chunk 两边都靠前，它会排得更高。",
            "Rerank（重排）: Cross-Encoder 逐条判断 query 和 chunk 是否真的相关。",
            "返回 answer（回答）: api_key 通常在 config/settings.yaml 的 llm 或 embedding 段配置。",
            "返回 evidence（证据）: citations 指向具体文档路径、chunk_id、标题或页码。",
        ],
    )

    add_heading(doc, "3. 核心技术栈逐个解释", 2)
    add_table(
        doc,
        ["技术", "通俗解释", "工作原理", "在项目中的例子"],
        [
            [
                "RAG",
                "让大模型先查资料，再回答。",
                "Retrieval 检索相关 chunk，把 chunk 作为上下文交给 LLM，减少瞎编。",
                "用户问配置问题，系统先从文档库找出相关片段，再生成带引用的回答。",
            ],
            [
                "ChromaDB",
                "存向量的数据库。",
                "每个 chunk 经过 embedding 变成高维向量；查询也变成向量，用相似度找最近 chunk。",
                "Embedding 后的文档块被写入 Chroma collection，查询时 DenseRetriever 读取。",
            ],
            [
                "BM25",
                "传统搜索引擎关键词打分算法。",
                "看查询词在文档里出现频率、逆文档频率和文档长度，适合代码名、配置名、专有名词。",
                "搜 `api_key`、`settings.yaml`、`BM25Indexer` 时，BM25 往往比纯向量更稳。",
            ],
            [
                "Dense Retrieval",
                "语义检索。",
                "把文本映射到向量空间，相似语义会靠近，即使词不一样也能召回。",
                "问“登录失败怎么办”，也可能召回“认证异常排查”这类表达不同但语义相近的文档。",
            ],
            [
                "RRF",
                "把多路排名合并的算法。",
                "公式是 `score = Σ 1 / (k + rank)`，只看排名，不依赖不同检索器的分数尺度。",
                "`src/core/query_engine/fusion.py` 的 `RRFFusion.fuse()` 把 dense 和 sparse 列表融合。",
            ],
            [
                "Cross-Encoder",
                "更慢但更准的精排模型。",
                "把 query 和候选 chunk 拼在一起输入模型，输出相关性分数。",
                "先召回 20 条，再 rerank 成最终 Top 5，提升答案相关度。",
            ],
            [
                "MCP Protocol",
                "让 AI 助手调用外部工具的协议。",
                "服务把能力注册成 tools，每个 tool 有 schema，助手按 schema 传参调用。",
                "`query_knowledge_hub` 输入 query/top_k/collection，输出回答和 citations。",
            ],
            [
                "Ragas / Custom Eval",
                "评估 RAG 回答质量。",
                "用 golden question 检查召回、相关性、忠实度，不靠肉眼感觉调参。",
                "面试可说：我用评测集看 hit rate、faithfulness、evidence coverage 是否退化。",
            ],
        ],
        [2.6, 3.5, 5.3, 5.6],
    )

    add_heading(doc, "4. 代码级链路：query_knowledge_hub 怎么跑", 2)
    add_numbers(
        doc,
        [
            "`QueryKnowledgeHubTool.execute(query, top_k, collection)` 先校验 query 是否为空，并确定 collection。",
            "`_ensure_initialized(collection)` 懒加载组件：EmbeddingFactory、VectorStoreFactory、BM25Indexer、DenseRetriever、SparseRetriever、HybridSearch、Reranker。",
            "`_perform_search()` 调用 HybridSearch，先取 `top_k * 2` 条候选，给后续 rerank 留空间。",
            "HybridSearch 内部做两路召回：dense_retriever 查 ChromaDB，sparse_retriever 查 BM25。",
            "`RRFFusion.fuse()` 把候选按排名融合，得到初始排序。",
            "`_apply_rerank()` 如果 reranker 可用，就用 Cross-Encoder 重排；失败时回退原排序。",
            "`ResponseBuilder.build()` 把结果组织成 Markdown 风格回答和 citations。",
            "TraceContext 记录 initialization、search、fusion、rerank 等阶段，Dashboard 可以展示全链路可观测信息。",
        ],
    )

    add_heading(doc, "5. 面试优先级问题", 2)
    add_table(
        doc,
        ["优先级", "面试官可能问", "推荐回答思路"],
        [
            ["P0", "为什么不用纯向量检索？", "纯向量适合理解语义，但对配置名、函数名、缩写、版本号不稳定；BM25 对关键词强，二者用 RRF 融合能兼顾查全率和查准率。"],
            ["P0", "RRF 为什么不用归一化分数？", "Dense 和 BM25 分数尺度不同，直接加权容易失真；RRF 只看排名，公式简单、稳定、对异构检索器友好。"],
            ["P0", "MCP 在项目里解决了什么？", "把 RAG 检索能力封装为标准工具，AI 助手不用知道内部数据库细节，只按 schema 调 `query_knowledge_hub`。"],
            ["P1", "摄取管线如何保证更新一致性？", "文档解析、chunk、metadata、embedding、BM25 都围绕 collection 和 source metadata 管理；更新时可删除旧 source 对应记录再重建。"],
            ["P1", "Rerank 的代价是什么？", "更准但更慢，所以先粗召回较少候选，再只对 Top N 做 Cross-Encoder 精排。"],
            ["P2", "如何评价 RAG 效果？", "准备 golden questions，看 hit rate、recall、faithfulness、answer relevancy、evidence coverage；线上看查询 trace 和失败样例。"],
        ],
        [1.5, 5.0, 10.5],
    )
    doc.add_page_break()


def build_project_two(doc: Document) -> None:
    add_heading(doc, "TwinMind Nexus（孪生智枢）项目面试准备手册", 1)
    add_para(
        doc,
        "这个项目是你自己的重点项目。它不是普通“上传文档问答”的 RAG，而是把一个完整代码项目变成 project digital twin（项目数字孪生）：上传 ZIP 后，系统自动扫描代码、文档、配置和图片，抽取实体、关系、证据卡，构建知识图谱，再让 Agent 像项目分析师一样规划任务、调用工具、验证证据、生成报告。",
    )

    add_callout(
        doc,
        "面试时一句话",
        "我做了一个多模态项目数字孪生平台 TwinMind Nexus，把 RAG 从“文档问答”升级成“代码项目理解”：系统会用 Tree-sitter 抽取多语言代码结构，用 Chroma + BM25 + RRF 做图谱增强 Hybrid RAG，再用 bounded ReAct Agent 自动规划探索路径、调用图谱和检索工具、验证证据链，并通过 AgentEval 做质量回归。",
        LIGHT_BLUE,
    )

    add_heading(doc, "1. 业务逻辑怎么讲", 2)
    add_table(
        doc,
        ["用户场景", "用户想做什么", "TwinMind 怎么做"],
        [
            ["看陌生项目", "上传一个 ZIP，快速知道入口、模块、依赖、配置和风险。", "生成项目档案：实体、关系、证据卡、展厅、图谱、报告。"],
            ["查架构关系", "不知道该搜什么，只想先看项目有哪些核心对象。", "图谱探索支持搜索、邻域展开、聚类、路径保存和实体详情。"],
            ["让 Agent 分析", "希望系统自己规划“先查什么、再查什么”。", "Planner 生成任务队列，ReAct task runner 调工具，Verifier 检查证据。"],
            ["可审计报告", "不想只看大模型总结，要知道依据来自哪里。", "finding 必须绑定 entity / relation / evidence ID，能回到证据卡。"],
            ["质量回归", "每次优化后想知道系统有没有变好。", "AgentEval 记录 evaluation_score、evidence_hit_rate、mission_verifier 等指标。"],
        ],
        [3.1, 5.2, 8.1],
    )

    add_heading(doc, "2. 前端具体怎么操作", 2)
    add_numbers(
        doc,
        [
            "运行 `python3 /Users/kitten/MultimodalRAG/start_twinmind_archive.py`，打开 React 工作台。",
            "进入“档案总览”，上传项目 ZIP，选择摄取范围：架构优先、完整审计、文档优先、测试质量分析。",
            "点击“生成档案”。后台会解压 ZIP、过滤重目录、扫描文件、抽取实体关系、构建图谱和 Hybrid RAG 索引。",
            "生成后看项目指标：实体数、关系数、证据数、展厅分布、Hybrid RAG 状态、多模态状态。",
            "进入“图谱探索”，搜索实体，例如 `PaymentService`、`BM25Indexer`、`settings.yaml`；点击节点查看邻居和证据链。",
            "进入“Agent 分析”，点击启动 ReAct 或运行 Agent 报告。这里会显示任务队列、工具调用轨迹、模型状态、报告和评测。",
            "进入“知识宇宙”，可以对多个项目档案做跨项目聚类、路径保存和对比探索。",
        ],
    )

    add_heading(doc, "3. 后端生成档案的完整链路", 2)
    add_code_like(
        doc,
        "从上传 ZIP 到生成档案：",
        [
            "前端上传 ZIP -> POST /api/archives/upload-job",
            "api.py:_run_upload_job -> 解压 ZIP、过滤 .git/node_modules/.venv/data 等目录",
            "ProjectArchiveService.ingest_project -> 调用 ArchiveBuilder.build",
            "ArchiveBuilder.build -> scanner.scan() 得到 ProjectFile 列表",
            "根据语言选择 adapter: PythonAdapter / MarkdownAdapter / ConfigAdapter / TreeSitterCodeAdapter / GenericAdapter",
            "adapter.extract -> 生成 entities（实体）、relations（关系）、evidence_cards（证据卡）",
            "enhance_semantic_graph -> 增强模块依赖、服务边界、配置依赖、RAG 组件等语义关系",
            "graph_store.upsert_* -> 写入 SQLite 图存储",
            "ProjectHybridRAGIndex.build -> 生成 evidence/entity/relation chunk，写入 Chroma + BM25",
            "MultiAgentPipeline 或 AgentMissionRuntime -> 生成 Agent 报告或 ReAct 任务轨迹",
        ],
    )

    add_heading(doc, "4. 核心数据结构怎么理解", 2)
    add_table(
        doc,
        ["概念", "中文解释", "项目例子"],
        [
            ["Entity（实体）", "项目里的一个对象。可以是文件、类、函数、接口、配置项、依赖、文档概念。", "`PaymentService` 是类实体；`settings.yaml` 是配置文件实体；`BM25Indexer` 是类实体。"],
            ["Relation（关系）", "两个实体之间的连接。", "`File A IMPORTS dependency`，`Class DECLARES Method`，`Service DEPENDS_ON Config`。"],
            ["EvidenceCard（证据卡）", "证明某个实体或关系存在的原始片段。", "代码第几行的 import、函数定义片段、README 标题、图片识别描述。"],
            ["Hall（展厅）", "把实体按用途分组，方便浏览。", "架构展厅、检索展厅、配置展厅、概念展厅、依赖展厅。"],
            ["ProjectArchiveDraft", "一个项目档案草稿，包含上述所有内容。", "上传 `mall.zip` 后生成 `data/project_archive/java__macrozheng__mall/draft_archive.json`。"],
        ],
        [3.0, 5.1, 8.3],
    )

    add_heading(doc, "5. Agent 设计：不是随便聊天，而是受控探索", 2)
    add_para(
        doc,
        "TwinMind 的 Agent 采用“Supervisor Plan-and-Execute + bounded ReAct”的范式。通俗讲，就是先让 Planner 把一个大目标拆成任务队列；每个任务内部再用 ReAct 思路选择工具、观察结果、补证据；最后由 Verifier 检查 finding 是否有证据支撑。",
    )
    add_table(
        doc,
        ["模块", "对应代码/函数", "作用"],
        [
            ["Planner（规划器）", "`plan_agent_mission()`", "根据目标生成任务队列，例如找入口、映射展厅、检查核心实体、收集证据、总结架构。"],
            ["Task Queue（任务队列）", "`AgentMissionTask`", "每个任务有 objective、allowed_tools、max_steps、status、input_entity_ids。"],
            ["Tool Use（工具调用）", "`AgentToolRegistry.execute()`", "只允许调用注册过的项目工具，避免 Agent 任意执行危险操作。"],
            ["ReAct Runner", "`AgentMissionRuntime._run_to_completion()`", "每个任务循环选择 action、执行工具、记录 observation，直到完成或预算用完。"],
            ["Critic Retry（批评重试）", "任务发现证据不足时 bounded retry", "如果 finding 不确定，优先补 `hybrid_search` 或 `get_evidence`，不是无限重试。"],
            ["Verifier（验证器）", "`EvidenceVerifier.verify_task()`", "检查 finding 里引用的 evidence_id 是否真的存在，输出 accepted / partial / uncertain。"],
            ["Memory（记忆）", "`agent_memory.json` + mission scoped memory", "保存重要实体、关系、证据、评测历史和建议，避免每次从零开始。"],
        ],
        [3.0, 5.0, 8.4],
    )

    add_heading(doc, "6. ReAct 任务例子：输入一个问题后发生什么", 2)
    add_code_like(
        doc,
        "例子：用户点击“启动 ReAct”，目标是“理解这个 Java 项目的架构”。",
        [
            "输入 goal（目标）: Understand project architecture / 理解项目架构",
            "Planner 生成任务 1: find_entry_points（找入口点）",
            "允许工具: graph_summary、list_halls、graph_search、hybrid_search",
            "LLM 或确定性 fallback 选择 action: graph_summary({project_id:'java__macrozheng__mall'})",
            "Observation（观察）: 图谱有 13127 个实体、若干关系、5 个展厅。",
            "任务 2: map_archive_halls（映射展厅）",
            "Action: list_halls -> 返回架构展厅、配置展厅、依赖展厅等实体数量。",
            "任务 3: inspect_core_entities（检查核心实体）",
            "Action: graph_search(query:'application main controller service') -> 找入口类、Controller、Service。",
            "任务 4: collect_architecture_evidence（收集架构证据）",
            "Action: hybrid_search(query:'Spring Boot entry point configuration dependency') -> 返回证据卡 ID。",
            "Verifier 检查: 每个 finding 是否绑定真实 evidence_id。",
            "输出 answer/report（回答/报告）: 架构摘要 + 核心模块 + 风险 + 引用证据。",
        ],
    )

    add_heading(doc, "7. Agent 工具有哪些", 2)
    add_table(
        doc,
        ["工具名", "中文含义", "输入", "输出"],
        [
            ["graph_summary", "图谱总览", "project_id", "实体数、关系数、展厅数等总体指标。"],
            ["list_halls", "列出展厅", "project_id", "每个展厅的名称、描述、实体数量。"],
            ["graph_search", "搜索图谱实体", "project_id + query + limit", "相关实体列表、entity_id、evidence_id。"],
            ["graph_neighborhood", "展开实体邻域", "project_id + focus_entity_id/hall_id + depth", "周围节点、关系、证据 ID。"],
            ["inspect_entity", "检查单个实体", "project_id + entity_id", "实体详情、相连关系、邻居、证据卡。"],
            ["hybrid_search", "Hybrid RAG 检索", "project_id + query + top_k", "Chroma + BM25 + RRF 融合后的 chunk 结果。"],
            ["get_evidence", "读取证据卡", "project_id + evidence_ids", "证据标题、来源路径、片段、关联实体。"],
            ["run_specialist_agent", "运行专家 Agent", "role + prior_agents", "档案员/制图师/侦探/怀疑者/策展人角色输出。"],
        ],
        [3.2, 3.0, 5.0, 5.0],
    )

    add_heading(doc, "8. 五个专家 Agent 怎么分工", 2)
    add_table(
        doc,
        ["Agent", "中文角色", "做什么", "为什么需要"],
        [
            ["archivist", "证据档案员", "盘点证据卡、入口路径、重要文件。", "先知道有哪些可靠材料。"],
            ["cartographer", "图谱制图师", "把实体、关系、展厅和跨展厅连接画出来。", "解决“项目结构在哪里”的问题。"],
            ["detective", "架构侦探", "追踪模块协作、依赖中心、影响路径。", "把静态图谱转成架构理解。"],
            ["skeptic", "证据怀疑者", "挑战弱证据、缺证据、稀疏图谱和过度推断。", "降低大模型胡说风险。"],
            ["curator", "报告策展人", "整合前面角色，生成项目 briefing。", "把碎片分析变成可读报告。"],
        ],
        [2.4, 3.0, 6.0, 5.0],
    )

    add_heading(doc, "9. 技术栈逐个解释", 2)
    add_table(
        doc,
        ["技术", "原理", "为什么用", "项目例子"],
        [
            ["ReAct Agent", "Reason + Act：模型根据观察选择工具，再根据工具结果继续行动。", "适合未知项目探索，因为用户不一定知道该搜什么。", "任务内选择 `graph_search`、`hybrid_search`、`inspect_entity`。"],
            ["Plan-and-Execute", "先全局规划，再逐任务执行。", "大项目不能让 Agent 漫游；需要任务边界、预算和可追踪状态。", "`plan_agent_mission()` 生成 5 类架构任务。"],
            ["Tool-Use", "把能力封装成工具，并用 schema 限制输入。", "让 Agent 使用项目真实数据，而不是靠模型记忆猜。", "`AgentToolRegistry` 只允许 8 个白名单工具。"],
            ["Critic / Verifier", "对输出进行证据检查和错误标记。", "项目分析必须可信，不能只有漂亮总结。", "`EvidenceVerifier` 检查 evidence_id 是否有效。"],
            ["Mission Memory", "保存任务中发现的实体、关系、证据 ID。", "后续任务可以接着用前面结果，形成探索路径。", "agent_memory.json 保存历史 harness run 和重要证据。"],
            ["AgentEval", "Agent 质量回归体系。", "优化后要有真实数值，而不是主观感觉。", "最新记录：evaluation_score 0.8583，evidence_hit_rate 0.9167，mission_verifier 1.0。"],
            ["DeepSeek V4-Pro", "LLM 负责增强报告、选择工具、生成结构化 JSON。", "中文能力、成本和可接入性适合个人项目。", "结构化调用使用 JSON mode，减少 JSON repair。"],
            ["Knowledge Graph", "用节点和边表达项目知识。", "代码项目天然是关系网络：文件、类、函数、依赖、配置。", "实体是类/函数/配置，关系是 IMPORTS/DEPENDS_ON/DECLARES 等。"],
            ["Graph-enhanced Hybrid RAG", "检索结果不只是文本 chunk，还扩展到实体邻域、关系和证据卡。", "单纯文本检索看不到架构连接，图谱能补多跳关系。", "命中一个 Service 后，可以展开 Controller、Repository、配置依赖。"],
            ["Tree-sitter", "增量语法解析器，把代码解析成语法树。", "比正则更懂多语言代码结构。", "Java/C++/TS/JS/Go/Rust 抽取类、函数、import、struct、enum。"],
            ["Ollama Vision", "本地视觉模型理解图片。", "项目里的架构图、截图不能只当文件名，要转成可检索文字。", "图片生成 image evidence card，caption 写入 Hybrid RAG chunk。"],
        ],
        [3.0, 4.7, 4.9, 4.8],
    )

    add_heading(doc, "10. 为什么没有直接用 LangGraph", 2)
    add_para(
        doc,
        "面试官可能会问：你既然做 Agent，为什么不用 LangGraph？你的回答要诚实：当前版本没有直接引入 LangGraph，而是采用了 LangGraph-style 的显式状态流。原因是第一版核心目标是把项目摄取、图谱、证据链和工具调用跑通，使用 dataclass + service + store 能减少框架复杂度，也更容易调试。",
    )
    add_table(
        doc,
        ["对比点", "当前 TwinMind 做法", "如果迁移到 LangGraph"],
        [
            ["状态", "`AgentMission`、`AgentMissionTask`、`AgentTraceEvent` 显式保存。", "用 LangGraph State 存 mission、task、observations。"],
            ["节点", "Planner、ToolRegistry、Verifier 是普通 Python 类。", "变成 plan / act / observe / verify / report 节点。"],
            ["边", "代码里按任务循环和状态判断推进。", "用条件边控制 retry、stop、fallback。"],
            ["优势", "轻量、可控、容易定位 bug。", "复杂 Agent 流程更清晰，适合后续多分支、多角色并发。"],
        ],
        [3.0, 6.2, 6.2],
    )

    add_heading(doc, "11. AgentEval 与 Harness Engineering", 2)
    add_para(
        doc,
        "Harness Engineering 可以理解为“给 Agent 建一个测试和裁判系统”。它不只跑一次 Agent，而是固定输入、记录输出、比较指标，判断这次改动有没有让系统退化。",
    )
    add_table(
        doc,
        ["指标", "中文含义", "怎么理解"],
        [
            ["evaluation_score", "综合评测分", "把检索、证据覆盖、图谱命中等合成一个分数。"],
            ["relation_hit_rate", "关系命中率", "问题需要关系时，系统有没有找对相关边。"],
            ["evidence_hit_rate", "证据命中率", "回答是否能找到足够 evidence card 支撑。"],
            ["mission_verifier", "任务验证分", "ReAct mission 的 finding 是否被证据支持。"],
            ["agent_json_repair_count", "JSON 修复次数", "DeepSeek 输出不是合法 JSON 时系统修复了几次，越低越稳。"],
            ["agent_trust", "Agent 可信度", "综合 unsupported claims、证据覆盖、fallback 等信号。"],
        ],
        [4.2, 4.6, 7.6],
    )
    add_callout(
        doc,
        "你的真实优化结果可以这样讲",
        "我不是只做 UI 演示，而是给 Agent 建了质量回归门禁。优化前 evaluation_score 是 0.5125，relation_hit_rate 是 0.0；接入图谱邻域扩展、证据 backfill、bounded retry 后，evaluation_score 提升到 0.8583，relation_hit_rate 提升到 0.8333，evidence_hit_rate 提升到 0.9167，mission_verifier 达到 1.0。",
        LIGHT_YELLOW,
    )

    add_heading(doc, "12. 面试优先级问题", 2)
    add_table(
        doc,
        ["优先级", "面试官可能问", "推荐回答思路"],
        [
            ["P0", "你的项目和普通 RAG 有什么区别？", "普通 RAG 多是文档问答；我的项目面向完整代码项目，先构建实体/关系/证据卡知识图谱，再用 Hybrid RAG 和 Agent 多步探索，输出可审计架构报告。"],
            ["P0", "Agent 真的做了什么，不是按钮假效果吗？", "Agent 有任务规划、工具白名单、trace event、证据验证和 mission store。点击启动后会创建 mission，逐任务调用 graph_summary、graph_search、hybrid_search 等工具，并记录 observation。"],
            ["P0", "如何防止 Agent 胡说？", "三层控制：工具只读项目档案；finding 必须引用 entity/relation/evidence ID；Verifier 检查证据是否存在，缺证据就标 uncertain 或触发 bounded retry。"],
            ["P0", "为什么要知识图谱？", "代码项目不是纯文本，核心是结构关系。图谱能表达类、函数、配置、依赖之间的边，支持邻域展开、多跳推理和证据追踪。"],
            ["P1", "大项目 10 万实体图谱怎么看？", "不一次性全画，而是搜索驱动、邻域展开、聚类和层级视图；图上显示当前上下文，详情和证据在侧栏展示。"],
            ["P1", "Tree-sitter 比正则好在哪里？", "正则只能匹配文本模式，容易误判；Tree-sitter 解析语法树，能稳定识别类、函数、import、interface、struct 等结构。"],
            ["P1", "DeepSeek 返回 JSON 不稳定怎么办？", "使用 provider JSON mode、temperature=0、短 schema、示例 JSON、max_tokens 限制；repair 只是兜底，并在 AgentEval 里统计 repair_count。"],
            ["P2", "为什么图片也进 RAG？", "项目截图、架构图、流程图包含重要信息。Vision 模型把图片转成 caption，再作为 image evidence card 进入 Chroma/BM25，可被搜索和引用。"],
            ["P2", "后续为什么可能迁移 LangGraph？", "当前自研状态流更轻；当任务分支、并行角色、复杂中断恢复变多时，可以把 Planner/Tool/Verifier 映射为 LangGraph 节点和条件边。"],
        ],
        [1.5, 5.0, 10.5],
    )
    doc.add_page_break()


def build_twinmind_scripts(doc: Document) -> None:
    add_heading(doc, "面试表达模板与追问速记", 1)
    add_para(
        doc,
        "这一部分专门帮你把 TwinMind Nexus 讲成面试语言。你可以先背 1 分钟版本，再根据面试官追问切到 3 分钟版本和技术细节。",
    )

    add_heading(doc, "1 分钟介绍模板", 2)
    add_callout(
        doc,
        "TwinMind Nexus 版",
        "我做的是 TwinMind Nexus，一个多模态项目数字孪生与自主 Agent 图谱 RAG 平台。用户上传完整项目 ZIP 后，系统会用 Tree-sitter 抽取多语言代码结构，生成实体、关系和证据卡知识图谱；再通过 Chroma + BM25 + RRF 做图谱增强 Hybrid RAG。Agent 部分采用 Plan-and-Execute + bounded ReAct，Planner 先拆任务，Task Runner 调用图谱和检索工具，Verifier 检查证据链，最后输出可审计项目报告。我还做了 AgentEval 质量回归，用 evaluation_score、evidence_hit_rate、mission_verifier 等指标验证优化效果。",
        LIGHT_BLUE,
    )

    add_heading(doc, "3 分钟深挖模板", 2)
    add_numbers(
        doc,
        [
            "先讲背景：普通 RAG 只能回答文档问题，但理解代码项目需要知道模块、接口、配置、依赖和证据。",
            "再讲输入输出：输入是 ZIP，输出是项目档案、图谱、证据卡、Agent 报告和评测结果。",
            "再讲摄取链路：上传 -> 解压过滤 -> scanner -> adapter -> Tree-sitter -> semantic_enhancer -> graph_store -> HybridRAGIndex。",
            "再讲 Agent：Planner 生成任务，ReAct 在每个任务里选择工具，ToolRegistry 限制工具，Verifier 绑定证据，Memory 记录重要结果。",
            "最后讲结果：evaluation_score 从 0.5125 到 0.8583，evidence_hit_rate 到 0.9167，mission_verifier 到 1.0。",
        ],
    )

    add_heading(doc, "最容易被追问的技术原理速记", 2)
    add_table(
        doc,
        ["关键词", "一句话原理", "你的项目里的落点"],
        [
            ["RAG", "先检索，再生成，减少模型幻觉。", "查询项目档案时先找证据卡和相关实体。"],
            ["Hybrid RAG", "语义向量 + 关键词 + 融合排序。", "Chroma 召回语义，BM25 匹配代码名，RRF 合并。"],
            ["Knowledge Graph", "节点表示对象，边表示关系。", "类/函数/配置是实体，IMPORTS/DEPENDS_ON 是关系。"],
            ["ReAct", "观察后选择工具，工具结果再影响下一步。", "Agent 根据 graph_summary 决定继续 graph_search 或 hybrid_search。"],
            ["Verifier", "检查结论有没有证据支持。", "finding 没有 evidence_id 就标 uncertain 或补证据。"],
            ["AgentEval", "给 Agent 建评测门禁。", "统计 JSON 稳定性、证据命中率、任务完成率。"],
            ["LangGraph", "把 Agent 流程建成状态图。", "当前未直接用，未来可把 Planner/Tool/Verifier 映射成节点。"],
        ],
        [3.0, 5.2, 8.2],
    )

    add_heading(doc, "背诵优先级", 2)
    add_bullets(
        doc,
        [
            "第一优先级：讲清楚 TwinMind Nexus 不是普通文档问答，而是项目数字孪生 + Agent 图谱 RAG。",
            "第二优先级：讲清楚 Agent 链路：Planner -> Tool Use -> Observation -> Critic Retry -> Verifier -> Memory。",
            "第三优先级：讲清楚证据链：answer（回答）不是凭空生成，必须能回到 evidence（证据卡）。",
            "第四优先级：讲清楚评测指标：为什么有 evaluation_score、evidence_hit_rate、mission_verifier、json_repair_count。",
            "第五优先级：讲清楚底层 RAG 基础：Chroma、BM25、RRF、证据卡、多模态图片 caption。",
        ],
    )


def build_doc() -> None:
    doc = setup_document()
    build_cover(doc)
    build_project_two(doc)
    build_twinmind_scripts(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
    print(OUTPUT)
